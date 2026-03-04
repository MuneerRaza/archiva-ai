import asyncio
import re
from pathlib import Path

from app.services.processing.extractors.base import Extractor, ExtractionResult, TableData
from app.services.processing.extractors.text import _detect_language


def _extract_docx_sync(file_path: Path) -> ExtractionResult:
    from docx import Document

    doc = Document(str(file_path))
    paragraphs: list[str] = []
    tables: list[TableData] = []
    headings: list[dict] = []

    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

        # Detect headings from paragraph style
        style_name = para.style.name if para.style else ""
        if style_name.startswith("Heading"):
            match = re.search(r"(\d+)$", style_name)
            level = int(match.group(1)) if match else 1
            title = para.text.strip()
            if title:
                headings.append({"title": title, "level": level})
        elif style_name == "Title":
            title = para.text.strip()
            if title:
                headings.insert(0, {"title": title, "level": 0})

    for table in doc.tables:
        rows_data: list[list[str]] = []
        for row in table.rows:
            rows_data.append([cell.text.strip() for cell in row.cells])

        if not rows_data:
            continue

        headers = rows_data[0]
        data_rows = rows_data[1:] if len(rows_data) > 1 else []
        tables.append(TableData(headers=headers, rows=data_rows))

    raw_text = "\n\n".join(paragraphs)
    words = raw_text.split()

    # Build structure from headings with anchor IDs for navigation
    structure = None
    if headings:
        doc_title = headings[0]["title"] if headings[0].get("level", 1) <= 1 else "Untitled"
        sections = []
        for i, h in enumerate(headings):
            sections.append({"title": h["title"], "level": h["level"], "anchor": f"heading-{i}"})
        structure = {
            "title": doc_title,
            "sections": sections,
        }

    return ExtractionResult(
        raw_text=raw_text,
        page_count=max(1, len(words) // 250),
        word_count=len(words),
        language=_detect_language(raw_text),
        tables=tables,
        structure=structure,
    )


class DocxExtractor(Extractor):
    async def extract(self, file_path: Path, mime_type: str) -> ExtractionResult:
        return await asyncio.to_thread(_extract_docx_sync, file_path)
