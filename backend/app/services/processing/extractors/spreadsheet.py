import asyncio
from pathlib import Path

from app.services.processing.extractors.base import Extractor, ExtractionResult, TableData


def _extract_spreadsheet_sync(file_path: Path) -> ExtractionResult:
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    tables: list[TableData] = []
    text_parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_data: list[list[str]] = []

        for row in ws.iter_rows(values_only=True):
            str_row = [str(cell) if cell is not None else "" for cell in row]
            # Skip completely empty rows
            if any(c for c in str_row):
                rows_data.append(str_row)

        if not rows_data:
            continue

        headers = rows_data[0]
        data_rows = rows_data[1:]

        tables.append(TableData(
            headers=headers,
            rows=data_rows,
            sheet_name=sheet_name,
        ))

        # Build text representation
        text_parts.append(f"Sheet: {sheet_name}")
        text_parts.append(" | ".join(headers))
        for row in data_rows[:50]:  # Limit text representation
            text_parts.append(" | ".join(row))
        if len(data_rows) > 50:
            text_parts.append(f"... ({len(data_rows) - 50} more rows)")
        text_parts.append("")

    wb.close()

    raw_text = "\n".join(text_parts)
    words = raw_text.split()

    return ExtractionResult(
        raw_text=raw_text,
        page_count=len(tables),
        word_count=len(words),
        language="en",
        tables=tables,
    )


class SpreadsheetExtractor(Extractor):
    async def extract(self, file_path: Path, mime_type: str) -> ExtractionResult:
        return await asyncio.to_thread(_extract_spreadsheet_sync, file_path)
